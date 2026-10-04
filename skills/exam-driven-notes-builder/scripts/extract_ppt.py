# -*- coding: utf-8 -*-
"""
PPT文本提取脚本（支持.pptx和.ppt）
功能：提取PPT文件中所有幻灯片的文本和表格内容
用法：python extract_ppt.py <PPT路径> <输出txt路径>
示例：python extract_ppt.py "教材/酶工程/5酶的分离纯化.ppt" "缓存/PPT_5酶的分离纯化.txt"

说明：
- .pptx文件用python-pptx库提取（pip install python-pptx）
- .ppt文件用win32com（PowerPoint COM自动化）提取，需安装PowerPoint
- 自动根据扩展名选择提取方式
"""
import os
import sys


def extract_pptx(path):
    """用python-pptx提取.pptx文本"""
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f'\n=== 幻灯片 {i} ===')
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        out.append(text)
            if shape.has_table:
                tbl = shape.table
                for row in tbl.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    out.append(' | '.join(cells))
    return '\n'.join(out)


def extract_ppt_via_com(path):
    """用win32com提取.ppt文本（需要安装PowerPoint）"""
    import win32com.client
    import pythoncom
    pythoncom.CoInitialize()
    ppt_app = None
    try:
        ppt_app = win32com.client.Dispatch("PowerPoint.Application")
        ppt_app.Visible = 1
        # 转为绝对路径，COM需要
        abs_path = os.path.abspath(path)
        pres = ppt_app.Presentations.Open(abs_path, WithWindow=False)
        out = []
        for i, slide in enumerate(pres.Slides, 1):
            out.append(f'\n=== 幻灯片 {i} ===')
            for shape in slide.Shapes:
                if shape.HasTextFrame:
                    if shape.TextFrame.HasText:
                        text = shape.TextFrame.TextRange.Text.strip()
                        if text:
                            out.append(text)
                if shape.HasTable:
                    tbl = shape.Table
                    for r in range(1, tbl.Rows.Count + 1):
                        cells = []
                        for c in range(1, tbl.Columns.Count + 1):
                            cells.append(tbl.Cell(r, c).Shape.TextFrame.TextRange.Text.strip())
                        out.append(' | '.join(cells))
        pres.Close()
        return '\n'.join(out)
    finally:
        if ppt_app:
            ppt_app.Quit()
        pythoncom.CoUninitialize()


def extract_ppt(path):
    """按扩展名自动选择提取方式，返回文本内容"""
    if path.lower().endswith('.pptx'):
        return extract_pptx(path)
    return extract_ppt_via_com(path)


def main():
    if len(sys.argv) < 3:
        print('用法: python extract_ppt.py <PPT路径> <输出txt路径>')
        sys.exit(1)

    ppt_path = sys.argv[1]
    out_path = sys.argv[2]

    print(f'处理: {ppt_path}')
    try:
        content = extract_ppt(ppt_path)
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(content)
        print(f'-> {out_path} ({len(content)} 字符)')
    except Exception as e:
        print(f'错误: {e}')
        sys.exit(1)


if __name__ == '__main__':
    main()
